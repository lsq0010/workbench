#!/usr/bin/env python3
"""多语言对照 —— 看翻译表、查漏、查占位符。

为什么做这个
  你桌面和下载夹里有 20 多个翻译表，最大的一份 4452 行 × 10 种语言
  （语言国际化.xlsx，中文/英文/法文/阿拉伯文/哈萨克语/俄语/孟加拉语/土耳其…）。
  这种表靠眼睛翻是翻不完的，而最容易出事的是两件：

    1. **占位符漏了** —— 中文有 【@SITE@】，某个语言的译文里没有，
       App 上就会显示成断掉的句子。这是真正会漏到线上的 bug。
    2. **漏翻** —— 某个语言的格子是空的。

  这两件事机器一秒能查完。所以这个工具的核心不是"看表"，
  而是**把这两类错挑出来**。

顺带能看表、搜内容、对比两个版本改了什么。

依赖 openpyxl（读 xlsx）。零网络。
"""
import json
import os
import re
import signal
import subprocess
import sys
import time
import urllib.parse
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

VERSION = "0.1.0"
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
HOME = os.path.abspath(os.environ.get("XLSCAT_HOME") or SCRIPT_DIR)
PIDFILE = os.path.join(HOME, "xlscat.pid")
LOGFILE = os.path.join(HOME, "xlscat.log")
REPORTS = os.path.join(HOME, "reports.jsonl")     # 只追加
DEFAULT_PORT = 8917

# 常见的占位符写法 —— 漏一个就可能让 App 显示断句
PLACEHOLDER_PATTERNS = [
    r"@[A-Za-z_][A-Za-z0-9_]*@",          # @SITE@
    r"\{[A-Za-z_0-9]+\}",                  # {name} {0}
    r"\{\{[^}]+\}\}",                      # {{name}}
    r"%[sdif@]",                            # %s %d
    r"\$\{[^}]+\}",                         # ${name}
    r"<[a-zA-Z/][^>]*>",                    # <b> </b>
    r"\[\[[^\]]+\]\]",                      # [[key]]
    r"\\n",                                 # 换行符
    r"\{[0-9]+\}",                          # {0} {1}
]


def log(msg):
    line = f"{datetime.now().strftime('%H:%M:%S')} {msg}"
    print(line, flush=True)
    try:
        with open(LOGFILE, "a", encoding="utf-8") as f:
            f.write(line + "\n")
    except Exception:
        pass


def _user_id():
    try:
        p = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "lib")
        if p not in sys.path:
            sys.path.insert(0, p)
        import platform_lib
        return platform_lib.user_id()
    except Exception:
        return ""


def have_openpyxl():
    try:
        import openpyxl  # noqa
        return True, getattr(openpyxl, "__version__", "?")
    except Exception as exc:
        return False, str(exc)


def placeholders(text):
    """把一段文案里的占位符全抠出来（归一化成集合）"""
    if not text:
        return set()
    out = set()
    for pat in PLACEHOLDER_PATTERNS:
        for m in re.findall(pat, str(text)):
            out.add(m)
    return out


# ══════════════════════════════════════════════════════════════
# 读表
# ══════════════════════════════════════════════════════════════

def sheet_names(path):
    if not have_openpyxl()[0]:
        return []
    import openpyxl
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
        names = wb.sheetnames
        wb.close()
        return names
    except Exception as exc:
        log(f"[sheet_names] {exc}")
        return []


def read_sheet(path, sheet=None, max_rows=20000):
    """读一个 sheet，返回表头和行"""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        import csv
        with open(path, encoding="utf-8-sig", errors="replace") as f:
            rows = [r for r in csv.reader(f)]
        sheet = sheet or "(csv)"
    else:
        if not have_openpyxl()[0]:
            return {"error": "没装 openpyxl，装一下：python3 -m pip install --user openpyxl"}
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if sheet not in wb.sheetnames:
            sheet = wb.sheetnames[0]
        ws = wb[sheet]
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                break
            rows.append(list(row))
        wb.close()
    # 补齐列数
    width = max((len(r) for r in rows), default=0)
    for r in rows:
        while len(r) < width:
            r.append(None)
    rows = [[("" if c is None else c) for c in r] for r in rows]
    return {"sheet": sheet, "rows": rows, "width": width, "height": len(rows)}


def guess_header(rows):
    """猜哪一行是语言表头 —— 通常是第一行非空且多为文本"""
    for i, r in enumerate(rows[:6]):
        nonempty = [c for c in r if str(c).strip()]
        if len(nonempty) >= 2 and all(len(str(c)) < 24 for c in nonempty):
            return i
    return 0


def analyze(path, sheet=None, key_col=0, header_row=None):
    """核心：查漏翻 + 查占位符不一致"""
    d = read_sheet(path, sheet)
    if d.get("error"):
        return d
    rows = d["rows"]
    if not rows:
        return {"error": "这个表是空的"}
    hdr_i = header_row if header_row is not None else guess_header(rows)
    header = [str(c).strip() for c in rows[hdr_i]]
    body = rows[hdr_i + 1:]
    ncol = len(header)

    # 语言列：表头非空的列（跳过完全空白的列）
    lang_cols = [i for i in range(ncol) if header[i]]

    issues = []          # 问题清单
    stats = {i: {"total": 0, "empty": 0, "ph_bad": 0} for i in lang_cols}
    checked_rows = 0

    for ri, r in enumerate(body):
        if not any(str(c).strip() for c in r):
            continue
        src = str(r[key_col]).strip() if key_col < len(r) else ""
        if not src:
            continue          # 没有原文的行跳过（可能是说明行）
        # 语言代码行（ZH-CN / en-US / zh_CN）不是待翻译文案，跳过
        if len(src) <= 12 and re.fullmatch(r"[A-Za-z]{2,3}[-_][A-Za-z0-9]{2,7}", src):
            continue
        checked_rows += 1
        src_ph = placeholders(src)
        row_label = src[:60]

        for ci in lang_cols:
            if ci == key_col:
                continue
            val = str(r[ci]).strip() if ci < len(r) else ""
            stats[ci]["total"] += 1
            if not val:
                stats[ci]["empty"] += 1
                issues.append({
                    "level": "warn", "kind": "漏翻",
                    "row": hdr_i + 2 + ri, "lang": header[ci],
                    "text": row_label,
                    "detail": "这个语言是空的",
                })
                continue
            val_ph = placeholders(val)
            missing = src_ph - val_ph
            extra = val_ph - src_ph
            if missing or extra:
                stats[ci]["ph_bad"] += 1
                bits = []
                if missing:
                    bits.append("少了 " + " ".join(sorted(missing)))
                if extra:
                    bits.append("多了 " + " ".join(sorted(extra)))
                issues.append({
                    "level": "critical", "kind": "占位符不一致",
                    "row": hdr_i + 2 + ri, "lang": header[ci], "col": ci,
                    "text": row_label,
                    "detail": "；".join(bits),
                    "value": val[:120],
                })

    # 同一个原文出现多次（可能翻译不一致）
    seen = {}
    for ri, r in enumerate(body):
        src = str(r[key_col]).strip() if key_col < len(r) else ""
        if not src:
            continue
        if src in seen:
            seen[src].append(hdr_i + 2 + ri)
        else:
            seen[src] = [hdr_i + 2 + ri]
    dup_src = {k: v for k, v in seen.items() if len(v) > 1}
    for src, lines in list(dup_src.items())[:200]:
        issues.append({
            "level": "info", "kind": "原文重复",
            "row": lines[0], "lang": header[key_col] if key_col < len(header) else "原文",
            "text": src[:60],
            "detail": "在第 %s 行重复出现" % "、".join(str(x) for x in lines[:6]),
        })

    # ── 行错位检测 ──
    # 翻译表最常见的严重错误不是"翻错词"，而是**整列错位**：
    # 复制粘贴时少了一行，后面全串位。症状是占位符对不上。
    # 判据：本行的**译文**占位符，恰好等于**邻近某行的原文**占位符
    # —— 说明这一格拿错了行。
    src_by_line = {}
    val_by_line_col = {}
    for ri, r in enumerate(body):
        ln = hdr_i + 2 + ri
        v = str(r[key_col]).strip() if key_col < len(r) else ""
        src_by_line[ln] = (v, placeholders(v))
        for ci in lang_cols:
            val_by_line_col[(ln, ci)] = str(r[ci]).strip() if ci < len(r) else ""

    misplaced = []
    for it in list(issues):
        if it["kind"] != "占位符不一致":
            continue
        ln = it["row"]
        ci = it.get("col")
        if ci is None:
            continue
        my_src = src_by_line.get(ln, ("", set()))[1]
        my_val = placeholders(val_by_line_col.get((ln, ci), ""))
        if my_val == my_src:
            continue                      # 占位符一致，不是错位
        # 只有译文里有**非空**占位符时才敢判断错位：
        # 两边都是空集合的话，任何一行都对得上，那是没证据，不是结论。
        if not my_val:
            continue
        # 在邻近 ±15 行里找：哪一行的原文占位符 == 本行译文的占位符
        for off in range(-15, 16):
            if off == 0:
                continue
            other = ln + off
            if other not in src_by_line:
                continue
            o_txt, o_ph = src_by_line[other]
            if o_ph == my_val and o_ph != my_src and o_txt:
                misplaced.append({
                    "level": "critical", "kind": "疑似行错位",
                    "row": ln, "lang": it["lang"], "col": ci, "text": it["text"],
                    "detail": "这一格的译文像是第 %d 行「%s」的（占位符完全对得上）"
                              % (other, o_txt[:26]),
                })
                break
    issues.extend(misplaced)

    order = {"critical": 0, "warn": 1, "info": 2}
    issues.sort(key=lambda x: (order.get(x["level"], 3), x["row"]))

    langs = []
    for ci in lang_cols:
        s = stats[ci]
        langs.append({
            "col": ci, "name": header[ci] or ("第 %d 列" % (ci + 1)),
            "total": s["total"], "empty": s["empty"], "ph_bad": s["ph_bad"],
            "done_pct": round((s["total"] - s["empty"]) / s["total"] * 100, 1)
                        if s["total"] else 100.0,
        })

    result = {
        "ok": True, "file": path, "sheet": d["sheet"], "header_row": hdr_i + 1,
        "header": header, "langs": langs, "rows_checked": checked_rows,
        "issues": issues[:2000], "issue_total": len(issues),
        "critical": len([x for x in issues if x["level"] == "critical"]),
        "warn": len([x for x in issues if x["level"] == "warn"]),
        "info": len([x for x in issues if x["level"] == "info"]),
        "dup_src": len(dup_src),
    }
    log(f"[analyze] {os.path.basename(path)}/{d['sheet']} {checked_rows} 行 "
        f"→ 占位符 {result['critical']} / 漏翻 {result['warn']}")
    return result


def page_rows(path, sheet=None, offset=0, limit=100, query=""):
    """翻页看表；query 非空时跨列搜索"""
    d = read_sheet(path, sheet, max_rows=20000)
    if d.get("error"):
        return d
    rows = d["rows"]
    hdr_i = guess_header(rows)
    header = [str(c) for c in rows[hdr_i]]
    body = [(i + hdr_i + 2, [str(c) for c in r]) for i, r in enumerate(rows[hdr_i + 1:])]
    if query:
        q = query.lower()
        body = [x for x in body if any(q in c.lower() for c in x[1])]
    total = len(body)
    page = body[offset:offset + limit]
    return {"ok": True, "header": header, "rows": page, "total": total,
            "offset": offset, "limit": limit, "sheet": d["sheet"],
            "sheets": sheet_names(path) if path.lower().endswith(".xlsx") else ["(csv)"]}


def diff_files(path_a, path_b, key_col=0, sheet=None):
    """对比两个版本的翻译表：多了什么、少了什么、改了什么"""
    a = read_sheet(path_a, sheet)
    b = read_sheet(path_b, sheet)
    if a.get("error"):
        return a
    if b.get("error"):
        return b

    def index(d):
        rows = d["rows"]
        hi = guess_header(rows)
        hdr = [str(c).strip() for c in rows[hi]]
        m = {}
        for r in rows[hi + 1:]:
            k = str(r[key_col]).strip() if key_col < len(r) else ""
            if k:
                m.setdefault(k, [str(c) for c in r])
        return hdr, m

    ha, ma = index(a)
    hb, mb = index(b)
    only_a = [k for k in ma if k not in mb]
    only_b = [k for k in mb if k not in ma]
    changed = []
    for k in ma:
        if k not in mb:
            continue
        ra, rb = ma[k], mb[k]
        for ci in range(min(len(ra), len(rb), len(hb))):
            if ci < len(ha) and ci < len(hb) and ha[ci] != hb[ci]:
                continue          # 表头列位置变了，跳过
            if ra[ci] != rb[ci]:
                changed.append({"key": k[:60], "col": hb[ci] if ci < len(hb) else str(ci),
                                "a": ra[ci][:90], "b": rb[ci][:90]})
                break
    return {"ok": True, "only_a": only_a[:300], "only_b": only_b[:300],
            "changed": changed[:500], "counts": {"只在A": len(only_a), "只在B": len(only_b),
                                                  "有改动": len(changed)}}


# ══════════════════════════════════════════════════════════════
# Web
# ══════════════════════════════════════════════════════════════

class Handler(BaseHTTPRequestHandler):
    protocol_version = "HTTP/1.1"
    server_version = "xlscat/" + VERSION

    def log_message(self, *a):
        pass

    def _send(self, code, body, ctype="application/json; charset=utf-8"):
        if isinstance(body, str):
            body = body.encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(body)

    def _body(self):
        n = int(self.headers.get("Content-Length") or 0)
        raw = self.rfile.read(n) if n else b""
        try:
            return json.loads(raw.decode("utf-8")) if raw else {}
        except Exception:
            return {}

    def do_GET(self):
        u = urllib.parse.urlsplit(self.path)
        qs = urllib.parse.parse_qs(u.query)
        if u.path in ("/", "/index.html"):
            try:
                with open(os.path.join(HOME, "index.html"), encoding="utf-8") as f:
                    return self._send(200, f.read(), "text/html; charset=utf-8")
            except Exception as exc:
                return self._send(500, f"<h1>{exc}</h1>", "text/html; charset=utf-8")

        if u.path == "/api/status":
            ok, ver = have_openpyxl()
            # 猜几个可能放表格的目录
            cands = [os.path.expanduser("~/Desktop/文档资料"),
                     os.path.expanduser("~/Desktop"),
                     os.path.expanduser("~/Downloads"),
                     os.path.expanduser("~/Desktop/周报")]
            return self._send(200, json.dumps({
                "version": VERSION, "openpyxl": ok, "openpyxl_version": ver,
                "folders": [c for c in cands if os.path.isdir(c)],
                "reports": reports_history(),
            }, ensure_ascii=False))

        if u.path == "/api/list":
            """列一个目录里的表格文件"""
            folder = (qs.get("folder", [os.path.expanduser("~/Desktop")])[0] or "").strip()
            out = []
            if os.path.isdir(folder):
                for name in sorted(os.listdir(folder)):
                    if name.startswith("."):
                        continue
                    p = os.path.join(folder, name)
                    if not os.path.isfile(p):
                        continue
                    if name.lower().endswith((".xlsx", ".xls", ".csv")):
                        try:
                            out.append({"name": name, "path": p,
                                        "bytes": os.path.getsize(p),
                                        "mtime": datetime.fromtimestamp(
                                            os.path.getmtime(p)).strftime("%Y-%m-%d %H:%M")})
                        except OSError:
                            pass
                out.sort(key=lambda x: -x["bytes"])
            return self._send(200, json.dumps({"folder": folder, "files": out[:200]},
                                              ensure_ascii=False))

        if u.path == "/api/sheets":
            p = (qs.get("path", [""])[0] or "").strip()
            if not os.path.isfile(p):
                return self._send(400, json.dumps({"error": "文件不存在"}, ensure_ascii=False))
            return self._send(200, json.dumps({"sheets": sheet_names(p)}, ensure_ascii=False))

        if u.path == "/api/digest":
            return self._send(200, json.dumps({"items": []}, ensure_ascii=False))
        return self._send(404, json.dumps({"error": "not found"}))

    def do_POST(self):
        u = urllib.parse.urlsplit(self.path)
        b = self._body()

        if u.path == "/api/analyze":
            p = (b.get("path") or "").strip()
            if not os.path.isfile(p):
                return self._send(400, json.dumps({"ok": False, "error": "文件不存在"},
                                                  ensure_ascii=False))
            r = analyze(p, b.get("sheet"), int(b.get("key_col") or 0),
                        b.get("header_row"))
            if r.get("ok"):
                with open(REPORTS, "a", encoding="utf-8") as f:      # 只追加
                    f.write(json.dumps({
                        "ts": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                        "user_id": _user_id(),
                        "file": os.path.basename(p), "sheet": r["sheet"],
                        "rows": r["rows_checked"], "critical": r["critical"],
                        "warn": r["warn"]}, ensure_ascii=False) + "\n")
            return self._send(200, json.dumps(r, ensure_ascii=False))

        if u.path == "/api/page":
            p = (b.get("path") or "").strip()
            if not os.path.isfile(p):
                return self._send(400, json.dumps({"ok": False, "error": "文件不存在"},
                                                  ensure_ascii=False))
            return self._send(200, json.dumps(
                page_rows(p, b.get("sheet"), int(b.get("offset") or 0),
                          int(b.get("limit") or 100), b.get("query") or ""),
                ensure_ascii=False))

        if u.path == "/api/diff":
            r = diff_files((b.get("a") or "").strip(), (b.get("b") or "").strip(),
                           int(b.get("key_col") or 0), b.get("sheet"))
            return self._send(200, json.dumps(r, ensure_ascii=False))

        if u.path == "/api/reveal":
            p = (b.get("path") or "").strip()
            target = p if os.path.isdir(p) else os.path.dirname(p)
            if target and os.path.isdir(target):
                subprocess.Popen(["open", target])
                return self._send(200, json.dumps({"ok": True}, ensure_ascii=False))
            return self._send(404, json.dumps({"ok": False, "error": "目录不存在"},
                                              ensure_ascii=False))
        return self._send(404, json.dumps({"error": "not found"}))


def reports_history(limit=10):
    rows = []
    try:
        with open(REPORTS, encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if line:
                    try:
                        rows.append(json.loads(line))
                    except Exception:
                        pass
    except Exception:
        pass
    return list(reversed(rows[-limit:]))


class Server(ThreadingHTTPServer):
    allow_reuse_address = True
    daemon_threads = True


def running_pid():
    try:
        pid = int((open(PIDFILE).read() or "0").strip())
    except Exception:
        return 0
    try:
        os.kill(pid, 0)
        return pid
    except Exception:
        return 0


def cmd_serve(args):
    port = args.port or DEFAULT_PORT
    with open(PIDFILE, "w") as f:
        f.write(str(os.getpid()))
    ok, ver = have_openpyxl()
    log(f"多语言对照 v{VERSION} 已启动 http://127.0.0.1:{port}"
        + (f"（openpyxl {ver}）" if ok else f"（⚠️ 没装 openpyxl：{ver}）"))

    def bye(*_):
        try:
            os.remove(PIDFILE)
        except Exception:
            pass
        sys.exit(0)

    signal.signal(signal.SIGTERM, bye)
    signal.signal(signal.SIGINT, bye)
    try:
        Server(("127.0.0.1", port), Handler).serve_forever()
    except OSError as exc:
        log(f"端口 {port} 起不来：{exc}")
        sys.exit(1)


def cmd_start(args):
    if running_pid():
        print(f"已在运行（PID {running_pid()}）")
        return 0
    port = args.port or DEFAULT_PORT
    logf = open(LOGFILE, "a")
    subprocess.Popen([sys.executable, os.path.realpath(__file__), "serve", "--port", str(port)],
                     stdout=logf, stderr=subprocess.STDOUT, stdin=subprocess.DEVNULL,
                     start_new_session=True, cwd=HOME)
    for _ in range(30):
        time.sleep(0.2)
        if running_pid():
            break
    print(f"✅ 多语言对照已启动（PID {running_pid()}） http://127.0.0.1:{port}"
          if running_pid() else "❌ 启动失败")
    return 0 if running_pid() else 1


def cmd_stop(args):
    pid = running_pid()
    if not pid:
        print("没有在运行")
        return 0
    try:
        os.kill(pid, signal.SIGTERM); time.sleep(0.6)
    except Exception:
        pass
    try:
        os.remove(PIDFILE)
    except Exception:
        pass
    print("已停止")
    return 0


def cmd_check(args):
    r = analyze(args.path, args.sheet, args.key_col)
    if r.get("error"):
        print("❌ " + r["error"])
        return 1
    print("  %s / %s　%d 行 × %d 种语言（表头在第 %d 行）"
          % (os.path.basename(r["file"]), r["sheet"], r["rows_checked"],
             len(r["langs"]), r["header_row"]))
    print()
    print("  ── 各语言完成度 ──")
    for l in r["langs"]:
        bar = "█" * int(l["done_pct"] / 5)
        flag = "" if l["done_pct"] >= 100 else ("  ← 漏 %d" % l["empty"])
        print("    %-14s %5.1f%% %-20s %s%s" % (l["name"][:14], l["done_pct"], bar,
                                                ("占位符错 %d" % l["ph_bad"]) if l["ph_bad"] else "",
                                                flag))
    print()
    print("  ── 问题：占位符 %d / 漏翻 %d / 重复原文 %d ──"
          % (r["critical"], r["warn"], r["info"]))
    for it in r["issues"][:25]:
        mark = {"critical": "❌", "warn": "⚠️", "info": "💡"}.get(it["level"], "·")
        print("    %s 第%-5d %-12s %s" % (mark, it["row"], it["lang"][:12], it["detail"]))
        print("         %s" % it["text"][:70])
    if r["issue_total"] > 25:
        print("    …… 还有 %d 条" % (r["issue_total"] - 25))
    return 0


def main():
    import argparse
    p = argparse.ArgumentParser(prog="xlscat", description=f"多语言对照 v{VERSION}")
    sub = p.add_subparsers(dest="cmd", required=True)
    sp = sub.add_parser("serve"); sp.add_argument("--port", type=int); sp.set_defaults(f=cmd_serve)
    sp = sub.add_parser("start"); sp.add_argument("--port", type=int); sp.set_defaults(f=cmd_start)
    sub.add_parser("stop").set_defaults(f=cmd_stop)
    sp = sub.add_parser("check"); sp.add_argument("path")
    sp.add_argument("--sheet"); sp.add_argument("--key-col", dest="key_col", type=int, default=0)
    sp.set_defaults(f=cmd_check)
    args = p.parse_args()
    sys.exit(args.f(args))


if __name__ == "__main__":
    main()
